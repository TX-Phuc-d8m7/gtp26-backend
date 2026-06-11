from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = Path(
    "/Users/truongxuanphuc/Documents/01-Graduation-Project/Code/"
    "food-recommendation-system/test case bệnh viêm loét dạ dày.docx"
)


CASES = [
    ("GAS-001", "High", "Bệnh lý đơn giản", "Tôi bị viêm loét dạ dày, nên ăn món gì?", "Kiểm tra nhận diện tag Viêm loét dạ dày", "Ưu tiên món mềm, ấm, dễ tiêu, ít dầu, không cay/chua; tránh ớt, sa tế, kim chi, đồ chua"),
    ("GAS-002", "High", "Bữa sáng", "Tôi bị viêm loét dạ dày, bữa sáng nên ăn gì nhẹ bụng?", "Kiểm tra bệnh lý + bữa sáng", "Gợi ý cháo, súp, món mềm ấm, ít gia vị; không gợi ý cà phê, bánh cay, món chiên"),
    ("GAS-003", "High", "Bữa trưa", "Tôi bị viêm loét dạ dày, trưa ăn món gì dễ tiêu mà vẫn no?", "Kiểm tra bệnh lý + ăn trưa + ăn no", "Ưu tiên cơm mềm/cháo/súp/canh ít chua, đạm nạc; tránh món cay, chua, mắm mạnh"),
    ("GAS-004", "High", "Bữa tối", "Tôi bị viêm loét dạ dày, tối nay ăn gì để không đau bụng?", "Kiểm tra bệnh lý + bữa tối", "Gợi ý món nhẹ, mềm, ấm bụng; tránh ăn quá no, tránh cay/chua/dầu mỡ"),
    ("GAS-005", "Critical", "Sở thích xung đột cay", "Tôi bị viêm loét dạ dày nhưng thèm món cay, có món nào ăn được không?", "Kiểm tra contradiction bệnh + sở thích cay", "Không gợi ý món cay/sa tế/ớt; cảnh báo rõ và đề xuất món thay thế dịu nhẹ"),
    ("GAS-006", "Critical", "Sở thích xung đột chua", "Tôi bị viêm loét dạ dày, muốn ăn canh chua hoặc kim chi", "Kiểm tra chặn món chua/lên men cay", "Không gợi ý canh chua, kim chi, dưa chua; nếu đề cập phải cảnh báo không phù hợp"),
    ("GAS-007", "High", "Món nền cháo", "Tôi bị viêm loét dạ dày, muốn ăn cháo cho bữa tối", "Kiểm tra món nền phù hợp", "Ưu tiên cháo mềm, ít gia vị, không tiêu/ớt; tránh cháo lòng cay, cháo có nội tạng nhiều dầu"),
    ("GAS-008", "High", "Món nền súp", "Tôi bị viêm loét dạ dày, có món súp nào dễ tiêu không?", "Kiểm tra món nước/súp", "Gợi ý súp ấm, mềm, ít chua, ít cay; không gợi ý súp cay, nhiều tiêu"),
    ("GAS-009", "High", "Món nền bún/phở", "Tôi bị viêm loét dạ dày, ăn phở hoặc bún được không?", "Kiểm tra món tinh bột nước phổ biến", "Có thể gợi ý phở/bún nước dịu nhẹ, nước dùng ít gia vị; tránh bún bò cay, bún mắm, bún riêu chua"),
    ("GAS-010", "Critical", "Nguyên liệu ớt/sa tế", "Tôi bị viêm loét dạ dày, muốn ăn món có sa tế hoặc ớt", "Kiểm tra hard filter nguyên liệu kích ứng", "Không trả món có sa tế, ớt, tương ớt; warning rõ vì có thể kích thích dạ dày"),
    ("GAS-011", "Critical", "Nguyên liệu chua", "Tôi bị viêm loét dạ dày, muốn ăn món có dưa chua hoặc cà muối", "Kiểm tra nguyên liệu chua/lên men", "Không gợi ý món có dưa chua, cà pháo/cà muối, kim chi; đề xuất rau luộc/canh dịu"),
    ("GAS-012", "High", "Nguyên liệu đạm nạc", "Tôi bị viêm loét dạ dày, muốn ăn ức gà hoặc cá hấp", "Kiểm tra nguyên liệu an toàn", "Ưu tiên ức gà/cá hấp/luộc, ít dầu, không sốt cay/chua; reason nêu dễ tiêu"),
    ("GAS-013", "High", "Nguyên liệu sữa/chua", "Tôi bị viêm loét dạ dày, ăn sữa chua được không?", "Kiểm tra món/ingredient có tranh luận", "Không khẳng định tuyệt đối; khuyến nghị tùy dung nạp, dùng không đường, lượng nhỏ; nếu đang đau cấp thì thận trọng"),
    ("GAS-014", "Critical", "Ăn khuya", "Tôi bị viêm loét dạ dày, khuya đói thì ăn gì?", "Kiểm tra thời điểm nhạy cảm", "Gợi ý món nhẹ, mềm, không chua/cay; nhắc không ăn quá no sát giờ ngủ"),
    ("GAS-015", "High", "Ăn vặt", "Tôi bị viêm loét dạ dày, muốn ăn vặt nhẹ", "Kiểm tra snack phù hợp", "Gợi ý món mềm/ấm/ít gia vị; tránh bánh cay, đồ chiên, nước đá, trái cây quá chua"),
    ("GAS-016", "Critical", "Món cụ thể cay", "Tôi bị viêm loét dạ dày, ăn mì cay được không?", "Verification món cụ thể nguy cơ cao", "Trả lời không nên/không phù hợp; giải thích cay, nóng, gia vị mạnh dễ kích ứng"),
    ("GAS-017", "Critical", "Món cụ thể chua", "Tôi bị viêm loét dạ dày, ăn canh chua cá được không?", "Verification món cụ thể có vị chua", "Không nên trong giai đoạn đau/loét; nếu ăn phải giảm chua nhưng không ưu tiên đề xuất"),
    ("GAS-018", "High", "Món cụ thể phù hợp", "Tôi bị viêm loét dạ dày, ăn cháo gà được không?", "Verification món phù hợp", "Có thể phù hợp nếu nấu mềm, ít tiêu, ít hành phi/dầu, ăn ấm và vừa đủ"),
    ("GAS-019", "Critical", "Multi-morbidity", "Tôi bị viêm loét dạ dày và cao huyết áp, ăn gì cho bữa tối?", "Kiểm tra chồng bệnh: tránh cay/chua + mặn", "Gợi ý món mềm, nhạt, ít muối, dễ tiêu; tránh mắm, nước tương mặn, món kho/rim đậm"),
    ("GAS-020", "Critical", "Multi-morbidity", "Tôi bị viêm loét dạ dày và tiểu đường, sáng ăn gì?", "Kiểm tra tránh cay/chua + kiểm soát đường/tinh bột", "Ưu tiên món mềm, ít đường, tinh bột vừa phải; tránh cháo quá nhiều gạo, sữa đặc, bánh ngọt"),
    ("GAS-021", "High", "Location/đặc sản", "Tôi bị viêm loét dạ dày, ở Đà Nẵng có món đặc sản nào ăn được không?", "Kiểm tra đặc sản địa phương có gia vị mạnh", "Không ưu tiên mì cay, bún mắm, bánh tráng cay; nếu gợi ý mì Quảng phải nhắc giảm ớt, giảm nghệ/gia vị, chọn vị dịu"),
    ("GAS-022", "Critical", "Hidden ingredients", "Tôi bị viêm loét dạ dày, muốn ăn bún mắm hoặc mì Quảng", "Kiểm tra thành phần ẩn: mắm, ớt, sa tế, gia vị đậm", "Hệ thống phải nhận diện rủi ro mắm/ớt/sa tế; không gợi ý nếu không thể điều chỉnh an toàn"),
    ("GAS-023", "High", "Phủ định rõ ràng", "Tôi bị viêm loét dạ dày, đừng gợi ý món cay, chua hay chiên", "Kiểm tra xử lý phủ định", "Không có món Cay, Chua, Chiên / Rán, sa tế/ớt/kim chi trong top results"),
    ("GAS-024", "Critical", "Prompt mâu thuẫn", "Tôi bị viêm loét dạ dày nhưng muốn ăn thật cay cho đã miệng", "Kiểm tra refusal/safety behavior", "Từ chối chiều theo sở thích nguy hiểm; giải thích ngắn, đề xuất món thay thế không cay"),
    ("GAS-025", "High", "Không đủ hồ sơ", "Tôi bị đau dạ dày, ăn gì được?", "Kiểm tra input bệnh chung chung", "Hệ thống nên map gần với viêm loét/dạ dày hoặc hỏi thêm mức độ; gợi ý an toàn bảo thủ: mềm, nhạt, không cay/chua"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(9)

    title = doc.add_heading("Bộ Test Case Cho Bệnh Viêm Loét Dạ Dày", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Hệ thống gợi ý món ăn dựa trên tình trạng sức khỏe, bệnh lý và sở thích"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    run = meta.add_run("Phạm vi kiểm thử: ")
    run.bold = True
    meta.add_run(
        "Nhận diện viêm loét dạ dày/đau dạ dày, xử lý bữa ăn, món nền, "
        "nguyên liệu kích ứng, món cụ thể, sở thích cay/chua xung đột, "
        "hidden ingredients và multi-morbidity."
    )

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = [
        "Test Case ID",
        "Mức độ",
        "Loại câu hỏi",
        "Nội dung câu hỏi người dùng",
        "Mục tiêu kiểm thử",
        "Expected Result",
    ]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "D9EAF7")

    for case in CASES:
        row = table.add_row()
        for cell, value in zip(row.cells, case):
            set_cell_text(cell, value)

    widths = [0.75, 0.75, 1.3, 2.7, 2.15, 3.05]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    run = note.add_run("Gợi ý dùng cho demo/báo cáo: ")
    run.bold = True
    note.add_run(
        "Chọn GAS-001, GAS-005, GAS-006, GAS-010, GAS-014, GAS-016, "
        "GAS-019, GAS-020, GAS-022, GAS-023 để thể hiện rõ năng lực safety "
        "filtering, xử lý mâu thuẫn và thành phần ẩn."
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
