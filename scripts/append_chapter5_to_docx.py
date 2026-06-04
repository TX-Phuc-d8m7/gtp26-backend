#!/usr/bin/env python3
"""Append the quantitative evaluation chapter to a DOCX report.

This script reads the JSON output produced by scripts/run_pathology_tests.py and
inserts a concise Chapter 5 before the KET LUAN section of a report copy.
"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK


def pct(value: float | int | None) -> str:
    if value is None:
        return "N/A"
    return f"{float(value):.2f}%"


def ms(value: float | int | None) -> str:
    if value is None:
        return "N/A"
    return f"{float(value):,.0f} ms"


def short(text: Any, limit: int = 180) -> str:
    value = str(text or "").strip()
    if len(value) <= limit:
        return value
    return value[: limit - 3].rstrip() + "..."


def get_group_rows(summary: dict[str, Any], key: str) -> list[dict[str, Any]]:
    value = summary.get(key) or []
    return value if isinstance(value, list) else []


def failure_cause(item: dict[str, Any]) -> str:
    reason = str(item.get("reason") or "").lower()
    if item.get("hard_filter_violation") and item.get("judge_status") == "PASS":
        return "False-positive của kiểm tra từ khóa"
    if item.get("pattern_group") == "Hidden Ingredient" and "cảnh báo" in reason:
        return "Thiếu cảnh báo thành phần ẩn"
    if any(
        word in reason
        for word in [
            "đậm đà",
            "ngọt",
            "tinh bột",
            "mồi nhậu",
            "chiên",
            "nướng",
            "chua",
            "cay",
            "purin",
            "nhiều dầu",
        ]
    ):
        return "Ràng buộc mềm/ranking chưa đủ mạnh"
    if any(
        word in reason
        for word in [
            "bơ",
            "chả cá",
            "pad thái",
            "bánh tráng trộn",
            "nước mắm",
            "mắm ruốc",
            "caffeine",
            "smoothie",
            "yogurt",
        ]
    ):
        return "Thiếu alias/dữ liệu thành phần"
    return "Khác"


def failure_cause_rows(results: list[dict[str, Any]]) -> list[list[str]]:
    labels = [
        "Ràng buộc mềm/ranking chưa đủ mạnh",
        "False-positive của kiểm tra từ khóa",
        "Thiếu cảnh báo thành phần ẩn",
        "Thiếu alias/dữ liệu thành phần",
        "Khác",
    ]
    counts = {label: 0 for label in labels}
    for item in results:
        if item.get("status") == "FAIL":
            counts[failure_cause(item)] = counts.get(failure_cause(item), 0) + 1
    return [[label, str(counts.get(label, 0))] for label in labels if counts.get(label, 0)]


def representative_failures(
    results: list[dict[str, Any]],
    max_rows: int,
) -> list[dict[str, Any]]:
    failed = [item for item in results if item.get("status") == "FAIL"]
    selected: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    for cause in [
        "Ràng buộc mềm/ranking chưa đủ mạnh",
        "Thiếu cảnh báo thành phần ẩn",
        "Thiếu alias/dữ liệu thành phần",
        "False-positive của kiểm tra từ khóa",
        "Khác",
    ]:
        for item in failed:
            if item.get("id") not in seen_ids and failure_cause(item) == cause:
                selected.append(item)
                seen_ids.add(item.get("id"))
                break

    for item in failed:
        if len(selected) >= max_rows:
            break
        if item.get("id") not in seen_ids:
            selected.append(item)
            seen_ids.add(item.get("id"))

    return selected[:max_rows]


def pass_rate(pass_count: int, total: int) -> float:
    return round(pass_count / total * 100, 1) if total else 0.0


def insert_paragraph(marker, text: str = "", style: str | None = None):
    return marker.insert_paragraph_before(text, style=style)


def insert_table(document: Document, marker, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    marker._p.addprevious(table._tbl)
    insert_paragraph(marker, "")


def find_conclusion_marker(document: Document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper() == "KẾT LUẬN":
            return paragraph
    raise ValueError("Khong tim thay muc 'KẾT LUẬN' de chen Chuong 5.")


def add_chapter(
    document: Document,
    marker,
    payload: dict[str, Any],
    max_fail_rows: int,
    smoke_results: list[dict[str, Any]] | None = None,
) -> None:
    summary = payload["summary"]
    results = payload.get("results") or []
    judge_pass = sum(1 for item in results if item.get("judge_status") == "PASS")
    judge_fail = len(results) - judge_pass
    programmatic_only_fail = sum(
        1
        for item in results
        if item.get("status") == "FAIL"
        and item.get("judge_status") == "PASS"
        and item.get("hard_filter_violation")
    )

    page_break = insert_paragraph(marker)
    page_break.add_run().add_break(WD_BREAK.PAGE)

    insert_paragraph(
        marker,
        "CHƯƠNG 5: ĐÁNH GIÁ ĐỊNH LƯỢNG VÀ THẢO LUẬN KẾT QUẢ",
        style="Heading 1",
    )
    insert_paragraph(
        marker,
        "Chương này trình bày kết quả đánh giá định lượng của hệ thống gợi ý món ăn theo ngữ cảnh sức khỏe. "
        "Mục tiêu của phần đánh giá là kiểm chứng khả năng nhận diện ràng buộc bệnh lý/dị ứng, mức độ an toàn "
        "của cơ chế lọc cứng, chất lượng xếp hạng sau truy xuất ngữ nghĩa và thời gian phản hồi của pipeline.",
    )

    insert_paragraph(marker, "5.1. Mục tiêu đánh giá", style="Heading 2")
    insert_paragraph(
        marker,
        "Đánh giá tập trung vào bốn câu hỏi chính: hệ thống có nhận diện đúng điều kiện sức khỏe trong câu hỏi "
        "ngôn ngữ tự nhiên hay không; các món có nguyên liệu chống chỉ định có bị loại khỏi kết quả hay không; "
        "kết quả cuối cùng có phù hợp với kỳ vọng của từng test case hay không; và pipeline có duy trì thời gian "
        "phản hồi chấp nhận được trong quá trình gọi LLM, embedding, lọc và xếp hạng hay không.",
    )

    insert_paragraph(marker, "5.2. Bộ test case sử dụng", style="Heading 2")
    insert_paragraph(
        marker,
        "Bộ benchmark chính gồm 35 test case bệnh lý, dị ứng, triệu chứng và tình huống tích hợp. Ngoài ra, "
        "đề tài bổ sung 60 test case mở rộng theo 4 pattern nhằm kiểm tra độ bền của hệ thống trước các cách "
        "diễn đạt tích cực, triệu chứng mơ hồ, ngữ cảnh bữa ăn và nguyên liệu ẩn. Kết quả được trình bày riêng "
        "cho benchmark chính và bộ mở rộng để không làm sai lệch phạm vi đánh giá đã mô tả trong phương pháp nghiên cứu.",
    )
    insert_table(
        document,
        marker,
        ["Bộ test", "Số case", "Mục đích"],
        [
            [
                "Benchmark chính",
                str(next((r.get("total", 0) for r in get_group_rows(summary, "by_source") if r.get("group") == "benchmark"), 35)),
                "Kiểm tra các tình huống cốt lõi về dị ứng, bệnh lý, triệu chứng và tích hợp nhiều ràng buộc.",
            ],
            [
                "Bộ mở rộng",
                str(next((r.get("total", 0) for r in get_group_rows(summary, "by_source") if r.get("group") == "extended"), 60)),
                "Kiểm tra độ bền theo bốn pattern: Positive Framing, Vague Symptom, Meal-Time Context và Hidden Ingredient.",
            ],
        ],
    )

    insert_paragraph(marker, "5.3. Phương pháp đánh giá", style="Heading 2")
    insert_paragraph(
        marker,
        "Mỗi test case được chạy trực tiếp qua hàm search_food() của backend nội bộ với cơ sở dữ liệu hiện tại. "
        "Kết quả trả về được đánh giá theo hai lớp. Lớp thứ nhất là kiểm tra tự động hard-filter/hard-rule, nhằm "
        "phát hiện món có nguyên liệu hoặc từ khóa bị cấm theo expected result. Lớp thứ hai là LLM-as-a-Judge, "
        "trong đó Gemini đánh giá kết quả thực tế dựa trên truy vấn, mục tiêu kiểm thử, expected result và danh sách "
        "món hệ thống gợi ý. Một test case chỉ được tính PASS khi không có vi phạm hard-filter và judge đánh giá đạt.",
    )

    insert_paragraph(marker, "5.4. Các chỉ số đánh giá", style="Heading 2")
    insert_paragraph(
        marker,
        "Các chỉ số được sử dụng gồm: số lượng PASS/FAIL tổng thể; PASS/FAIL theo benchmark và extended; PASS/FAIL "
        "theo nhóm dị ứng, bệnh lý, triệu chứng/tình trạng và tích hợp; PASS/FAIL theo bốn pattern mở rộng; số lỗi "
        "hard-filter; và thống kê latency gồm trung bình, nhỏ nhất, lớn nhất, p50 và p95.",
    )

    insert_paragraph(marker, "5.5. Kết quả định lượng", style="Heading 2")
    insert_paragraph(
        marker,
        f"Tổng cộng hệ thống được đánh giá trên {summary.get('total', 0)} test case. Kết quả chung đạt "
        f"{summary.get('pass', 0)} PASS, {summary.get('fail', 0)} FAIL, tương ứng tỷ lệ đạt {pct(summary.get('pass_rate'))}. "
        f"Số lỗi hard-filter được ghi nhận là {summary.get('hard_filter_violations', 0)}.",
    )
    insert_paragraph(
        marker,
        f"Đây là kết quả theo tiêu chí nghiêm ngặt: một case bị tính FAIL nếu không đạt LLM-as-a-Judge hoặc bị lớp "
        f"kiểm tra hard-filter tự động phát hiện vi phạm. Khi đối chiếu riêng lớp LLM-as-a-Judge, hệ thống đạt "
        f"{judge_pass}/{len(results)} case ({pct(pass_rate(judge_pass, len(results)))}). Có {programmatic_only_fail} case "
        f"bị FAIL chỉ do kiểm tra từ khóa tự động mặc dù judge đánh giá PASS; các case này được giữ trong thống kê "
        f"nghiêm ngặt nhưng được phân tích riêng như hạn chế của bộ kiểm thử.",
    )
    insert_table(
        document,
        marker,
        ["Bộ test", "Tổng số", "PASS", "FAIL", "Tỷ lệ PASS"],
        [
            [
                str(row.get("group", "")),
                str(row.get("total", 0)),
                str(row.get("pass", 0)),
                str(row.get("fail", 0)),
                pct(row.get("pass_rate")),
            ]
            for row in get_group_rows(summary, "by_source")
        ],
    )
    latency = summary.get("latency") or {}
    insert_table(
        document,
        marker,
        ["Chỉ số latency", "Giá trị"],
        [
            ["Trung bình", ms(latency.get("avg_ms"))],
            ["Nhỏ nhất", ms(latency.get("min_ms"))],
            ["Lớn nhất", ms(latency.get("max_ms"))],
            ["p50", ms(latency.get("p50_ms"))],
            ["p95", ms(latency.get("p95_ms"))],
        ],
    )
    insert_table(
        document,
        marker,
        ["Góc nhìn đánh giá", "PASS", "FAIL", "Tỷ lệ PASS"],
        [
            [
                "Kết quả nghiêm ngặt (hard-check + LLM-as-a-Judge)",
                str(summary.get("pass", 0)),
                str(summary.get("fail", 0)),
                pct(summary.get("pass_rate")),
            ],
            [
                "Đối chiếu riêng LLM-as-a-Judge",
                str(judge_pass),
                str(judge_fail),
                pct(pass_rate(judge_pass, len(results))),
            ],
        ],
    )
    if smoke_results:
        insert_paragraph(
            marker,
            "Ngoài đánh giá định lượng trên backend nội bộ, đề tài thực hiện smoke test bản triển khai tại "
            "https://gtp26-frontend.vercel.app/ thông qua proxy API của frontend. Smoke test chỉ nhằm xác nhận "
            "hệ thống deploy đang phản hồi được với các truy vấn đại diện, không được dùng thay thế cho metric chính.",
        )
        insert_table(
            document,
            marker,
            ["Nhóm truy vấn", "HTTP", "Latency", "Số kết quả", "Món trả về"],
            [
                [
                    str(item.get("group", "")),
                    str(item.get("http_status", "")),
                    ms(item.get("latency_ms")),
                    str(item.get("result_count", "")),
                    short(", ".join(item.get("results") or []), 120),
                ]
                for item in smoke_results
            ],
        )

    insert_paragraph(marker, "5.6. Phân tích kết quả theo nhóm", style="Heading 2")
    insert_paragraph(
        marker,
        "Kết quả theo nhóm cho phép quan sát nhóm ràng buộc nào hệ thống xử lý ổn định hơn và nhóm nào còn nhạy "
        "với dữ liệu hoặc cách diễn đạt truy vấn. Nhóm dị ứng thường yêu cầu độ an toàn cao nhất vì cần loại trừ "
        "nguyên liệu tuyệt đối; nhóm bệnh lý và triệu chứng có nhiều ràng buộc mềm nên kết quả phụ thuộc nhiều hơn "
        "vào scoring, penalty và cảnh báo.",
    )
    insert_table(
        document,
        marker,
        ["Nhóm", "Tổng số", "PASS", "FAIL", "Tỷ lệ PASS"],
        [
            [
                str(row.get("group", "")),
                str(row.get("total", 0)),
                str(row.get("pass", 0)),
                str(row.get("fail", 0)),
                pct(row.get("pass_rate")),
            ]
            for row in get_group_rows(summary, "by_category")
        ],
    )
    insert_table(
        document,
        marker,
        ["Pattern mở rộng", "Tổng số", "PASS", "FAIL", "Tỷ lệ PASS"],
        [
            [
                str(row.get("group", "")),
                str(row.get("total", 0)),
                str(row.get("pass", 0)),
                str(row.get("fail", 0)),
                pct(row.get("pass_rate")),
            ]
            for row in get_group_rows(summary, "by_pattern")
        ],
    )
    insert_paragraph(
        marker,
        "Nhìn chung, nhóm dị ứng đạt kết quả tốt hơn nhóm bệnh lý vì phần lớn dị ứng được xử lý bằng ràng buộc cứng "
        "theo nguyên liệu. Ngược lại, nhóm bệnh lý và tình huống tích hợp phụ thuộc nhiều vào ràng buộc mềm, điểm phạt "
        "tag và giải thích sau truy xuất nên dễ xuất hiện món còn tag cần hạn chế như Đậm đà, Ngọt, Giàu tinh bột hoặc "
        "Mồi nhậu. Trong bộ mở rộng, pattern Vague Symptom đạt cao nhất, cho thấy classifier có khả năng ánh xạ triệu "
        "chứng mơ hồ tương đối tốt. Pattern Hidden Ingredient thấp nhất do hệ thống chưa luôn trả lời trực diện về "
        "thành phần ẩn mà người dùng hỏi cụ thể.",
    )

    insert_paragraph(marker, "5.7. Phân tích các trường hợp chưa đạt", style="Heading 2")
    cause_rows = failure_cause_rows(results)
    if cause_rows:
        insert_paragraph(
            marker,
            "Các case chưa đạt được phân loại theo nguyên nhân chính để xác định hướng cải tiến. Bảng phân loại dưới "
            "đây dựa trên lý do đánh giá của LLM-as-a-Judge kết hợp cờ hard-filter của runner, dùng cho thảo luận và "
            "không làm thay đổi số liệu PASS/FAIL nghiêm ngặt.",
        )
        insert_table(
            document,
            marker,
            ["Nhóm nguyên nhân", "Số case"],
            cause_rows,
        )
    fail_rows = representative_failures(results, max_fail_rows)
    if fail_rows:
        insert_paragraph(
            marker,
            "Bảng dưới đây liệt kê các case FAIL tiêu biểu để phục vụ phân tích nguyên nhân. Các lỗi được ghi nhận "
            "trung thực thay vì loại khỏi thống kê, vì đây là căn cứ để xác định phạm vi cải tiến tiếp theo của hệ thống.",
        )
        insert_table(
            document,
            marker,
            ["ID", "Nguyên nhân", "Pattern", "Truy vấn", "Món trả về", "Lý do chưa đạt"],
            [
                [
                    str(item.get("id", "")),
                    failure_cause(item),
                    str(item.get("pattern_group") or "-"),
                    short(item.get("query"), 110),
                    short(item.get("actual_dishes_text"), 90),
                    short(item.get("reason"), 160),
                ]
                for item in fail_rows
            ],
        )
    else:
        insert_paragraph(marker, "Không ghi nhận test case FAIL trong lần chạy đánh giá này.")

    insert_paragraph(marker, "5.8. Hạn chế của phương pháp đánh giá", style="Heading 2")
    insert_paragraph(
        marker,
        "Phương pháp đánh giá vẫn có một số hạn chế. Thứ nhất, LLM-as-a-Judge phụ thuộc vào khả năng diễn giải của "
        "mô hình đánh giá nên có thể xuất hiện sai lệch trong các expected result quá nghiêm ngặt hoặc mơ hồ. Thứ hai, "
        "kiểm tra hard-filter theo từ khóa có thể tạo false positive đối với tên món chứa chuỗi nhạy cảm nhưng không "
        "thực sự chứa nguyên liệu nguy hiểm, ví dụ các biến thể món chay. Thứ ba, bộ dữ liệu món ăn còn hữu hạn nên "
        "một số yêu cầu đặc thù có thể không có món đáp ứng hoàn toàn. Cuối cùng, latency phụ thuộc vào Gemini, truy cập "
        "cơ sở dữ liệu và điều kiện mạng tại thời điểm chạy thử nghiệm.",
    )
    insert_paragraph(
        marker,
        "Từ kết quả trên, các hướng cải tiến cần ưu tiên gồm: chuẩn hóa lại checker hard-filter bằng ingredient key thay "
        "vì so khớp từ khóa thô; tăng penalty hoặc chuyển một số tag nguy cơ cao thành hard-filter trong các bệnh lý như "
        "tiểu đường, gout và tim mạch; bổ sung alias/thành phần mặc định cho các món có nguyên liệu ẩn như bánh tráng trộn, "
        "bơ trong bánh, mắm ruốc, dầu hào, nước mắm; và thêm bước trả lời trực diện khi truy vấn hỏi một thành phần cụ thể "
        "có an toàn hay không.",
    )
    insert_paragraph(
        marker,
        "Tuy vậy, bộ đánh giá 95 case vẫn cung cấp bằng chứng định lượng cần thiết cho báo cáo: hệ thống được kiểm thử "
        "trên benchmark chính đúng với phương pháp nghiên cứu và được kiểm tra bổ sung bằng các pattern mở rộng phản ánh "
        "cách người dùng thực tế đặt câu hỏi.",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-docx", required=True, type=Path)
    parser.add_argument("--results-json", required=True, type=Path)
    parser.add_argument("--smoke-json", type=Path)
    parser.add_argument("--output-docx", required=True, type=Path)
    parser.add_argument("--max-fail-rows", type=int, default=10)
    args = parser.parse_args()

    payload = json.loads(args.results_json.read_text(encoding="utf-8"))
    smoke_results = None
    if args.smoke_json:
        smoke_results = json.loads(args.smoke_json.read_text(encoding="utf-8"))
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source_docx, args.output_docx)

    document = Document(args.output_docx)
    marker = find_conclusion_marker(document)
    add_chapter(
        document,
        marker,
        payload,
        max_fail_rows=args.max_fail_rows,
        smoke_results=smoke_results,
    )
    document.save(args.output_docx)
    print(args.output_docx)


if __name__ == "__main__":
    main()
